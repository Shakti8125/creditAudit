"""Synthetic PD validation reports for the local demo (L-01).

Both reports describe the same fictional retail PD scorecard. Version 1 passes every
threshold in ``app/services/analytics/policy_checker.py``; version 2 breaches them, so the
demo shows a PASS-to-BREACH model lineage. Every name, address and number is fictional.
The masking demo needs something to mask, so the sign-off block carries a fictional
institution, two fictional people and a fictional e-mail address.

``python -m scripts.demo.fixtures OUT_DIR`` writes the two files, so a presenter can upload
one live. ``tests/test_demo_fixtures.py`` runs them through the real extractor and policy
checker, so a threshold change cannot silently break the demo.
"""

from __future__ import annotations

import argparse
import io
import math
from dataclasses import dataclass
from pathlib import Path

from docx import Document

DECILES = 10


@dataclass(frozen=True)
class ReportSpec:
    """One synthetic validation report."""

    key: str
    version_label: str
    filename: str
    expected_status: str
    gini: float  # percent, validation sample
    auc: float
    ks: float  # percent, validation sample
    psi: float
    hosmer_lemeshow_p: float
    brier: float
    verdict: str
    findings: tuple[str, ...]


REPORTS: tuple[ReportSpec, ...] = (
    ReportSpec(
        key="v1",
        version_label="1.0",
        filename="synthetic_retail_pd_validation_v1.docx",
        expected_status="PASS",
        gini=56.0,
        auc=0.78,
        ks=41.0,
        psi=0.06,
        hosmer_lemeshow_p=0.31,
        brier=0.09,
        verdict="The model is fit for purpose. No material finding was raised.",
        findings=(
            (
                "No material findings. Discrimination, calibration and stability are within "
                "the policy thresholds."
            ),
            (
                "Recommendation: continue annual monitoring and re-run this validation after "
                "any change to the scorecard variables."
            ),
        ),
    ),
    ReportSpec(
        key="v2",
        version_label="2.0",
        filename="synthetic_retail_pd_validation_v2.docx",
        expected_status="BREACH",
        gini=32.0,
        auc=0.66,
        ks=24.0,
        psi=0.31,
        hosmer_lemeshow_p=0.02,
        brier=0.28,
        verdict="The model is not fit for purpose in its current form. Remediation is required.",
        findings=(
            (
                "High finding: discriminatory power fell below the policy minimum. The "
                "scorecard no longer separates good and bad accounts adequately."
            ),
            (
                "High finding: the score distribution has drifted away from the development "
                "population, and the calibration test rejects the fitted probabilities."
            ),
            (
                "Recommendation: recalibrate or redevelop the scorecard, and report back to "
                "the model risk committee within one quarter."
            ),
        ),
    ),
)

REPORTS_BY_KEY = {spec.key: spec for spec in REPORTS}


def decile_table(target_psi: float) -> tuple[list[float], float]:
    """Returns the actual share per decile (percent) and the PSI of that table.

    The expected share is 10 % in every decile. The actual shares tilt linearly toward the
    high-risk deciles until the table's own PSI is ``target_psi`` to two decimals, so the PSI
    a report states always agrees with the table printed next to it.
    """
    shape = [-1 + 2 * i / (DECILES - 1) for i in range(DECILES)]

    def table(k: float) -> list[float]:
        actual = [round(10.0 * (1 + k * s), 1) for s in shape]
        # Rounding leaves the column a tenth short or over; put the residual in the middle.
        actual[DECILES // 2] = round(actual[DECILES // 2] + (100.0 - sum(actual)), 1)
        return actual

    def psi_of(actual: list[float]) -> float:
        return sum((a - 10.0) / 100 * math.log(a / 10.0) for a in actual)

    lo, hi = 0.0, 0.95
    for _ in range(60):
        mid = (lo + hi) / 2
        if psi_of(table(mid)) < target_psi:
            lo = mid
        else:
            hi = mid
    actual = table(hi)
    return actual, round(psi_of(actual), 2)


def _add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = text
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text


def build_report(spec: ReportSpec) -> bytes:
    """Builds one synthetic validation report as DOCX bytes."""
    actual, psi = decile_table(spec.psi)
    dev_gini = spec.gini + 2.0
    dev_auc = round(spec.auc + 0.01, 2)
    dev_ks = spec.ks + 1.5

    doc = Document()
    doc.core_properties.title = f"Synthetic retail PD scorecard validation, version {spec.version_label}"
    doc.core_properties.author = "Synthetic demo fixture"

    doc.add_heading(
        f"Retail PD Scorecard: Annual Model Validation Report (version {spec.version_label})", 0
    )
    doc.add_paragraph(
        "Synthetic document for demonstration. All names and figures are fictional."
    )

    doc.add_heading("1. Executive summary", 1)
    doc.add_paragraph(
        "On the out-of-time validation sample the model shows a Gini coefficient of "
        f"{spec.gini:.1f}%, an AUC of {spec.auc:.2f} and a KS statistic of {spec.ks:.1f}%. "
        f"The Population Stability Index (PSI) is {psi:.2f}. The Hosmer-Lemeshow p-value is "
        f"{spec.hosmer_lemeshow_p:.2f} and the Brier score is {spec.brier:.2f}."
    )
    doc.add_paragraph(spec.verdict)

    doc.add_heading("2. Scope and model description", 1)
    doc.add_paragraph(
        "The Retail PD Scorecard estimates the 12-month probability of default for unsecured "
        "personal-loan customers. It is a logistic regression with 14 variables built from "
        "application and behavioural data. A default is 90 days past due or an unlikely-to-pay "
        "flag. The development window runs from January 2021 to December 2023 and the "
        "out-of-time validation window from January 2024 to June 2025."
    )

    doc.add_heading("3. Discriminatory power", 1)
    doc.add_paragraph("Table 1 compares the development and validation samples.")
    _add_table(
        doc,
        ["Sample", "Gini", "AUC", "KS"],
        [
            ["Development", f"{dev_gini:.1f}%", f"{dev_auc:.2f}", f"{dev_ks:.1f}%"],
            ["Validation (out-of-time)", f"{spec.gini:.1f}%", f"{spec.auc:.2f}", f"{spec.ks:.1f}%"],
        ],
    )

    doc.add_heading("4. Calibration", 1)
    doc.add_paragraph(
        f"The Hosmer-Lemeshow p-value is {spec.hosmer_lemeshow_p:.2f}, and the Brier score is "
        f"{spec.brier:.2f}. Predicted and observed default rates were compared across ten "
        "score bands on the validation sample."
    )

    doc.add_heading("5. Population stability", 1)
    doc.add_paragraph(
        f"Table 2 compares the validation score distribution with the development "
        f"distribution. The Population Stability Index is {psi:.2f}."
    )
    _add_table(
        doc,
        ["Decile", "Expected %", "Actual %"],
        [
            [f"{i + 1}" + (" (lowest risk)" if i == 0 else " (highest risk)" if i == DECILES - 1 else ""),
             "10.0", f"{share:.1f}"]
            for i, share in enumerate(actual)
        ],
    )

    doc.add_heading("6. Findings and recommendations", 1)
    for finding in spec.findings:
        doc.add_paragraph(finding)

    doc.add_heading("7. Sign-off", 1)
    doc.add_paragraph(
        "Institution: Example National Bank. Prepared by Amara Okafor, Model Validation Unit. "
        "Reviewed by Daniel Whitfield, Head of Model Risk. "
        "Contact: model.validation@example.com."
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def write_fixtures(out_dir: Path) -> list[Path]:
    """Writes both reports into ``out_dir`` and returns their paths."""
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for spec in REPORTS:
        path = out_dir / spec.filename
        path.write_bytes(build_report(spec))
        paths.append(path)
    return paths


def main() -> None:
    parser = argparse.ArgumentParser(description="Write the synthetic demo reports.")
    parser.add_argument("out_dir", type=Path, help="directory to write the two DOCX files to")
    args = parser.parse_args()
    for path in write_fixtures(args.out_dir):
        print(path)


if __name__ == "__main__":
    main()
